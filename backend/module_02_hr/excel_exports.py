from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="000000")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
TITLE_FILL = "B8CCE4"
HEADER_FILL = "FFFF00"


def _sheet(title):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    return wb, ws


def _write_title_and_headers(ws, title, headers, col_widths):
    n_cols = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.fill = PatternFill("solid", fgColor=TITLE_FILL)
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.font = Font(bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


# openpyxl refuses timezone-aware datetimes outright — Excel has no
# timezone concept, so this renders the *local* wall-clock time as plain
# text rather than a raw (and rejected) tz-aware datetime object.
def _local_time(dt):
    return timezone.localtime(dt).strftime("%H:%M") if dt else None


def _finish(wb):
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


ATTENDANCE_HEADERS = ["Employee", "Department", "Date", "Status", "Check in", "Check out", "Hours"]
ATTENDANCE_WIDTHS = [22, 18, 12, 12, 10, 10, 8]
ATTENDANCE_STATUS_COLOR = {
    "PRESENT": "C6EFCE",
    "HALF_DAY": "FFEB9C",
    "ON_LEAVE": "DCE6F1",
    "ABSENT": "FFC7CE",
    "NO_LOGIN": "F2F2F2",
}
ATTENDANCE_STATUS_LABEL = {
    "PRESENT": "Present",
    "HALF_DAY": "Half day",
    "ON_LEAVE": "On leave",
    "ABSENT": "Absent",
    "NO_LOGIN": "No login",
}


def build_attendance_excel(rows, title):
    wb, ws = _sheet("Attendance")
    _write_title_and_headers(ws, title, ATTENDANCE_HEADERS, ATTENDANCE_WIDTHS)

    for row_idx, r in enumerate(rows, start=3):
        values = [
            r["full_name"],
            r["department_name"] or "—",
            r["date"],
            ATTENDANCE_STATUS_LABEL.get(r["status"], r["status"]),
            _local_time(r["check_in_time"]) or "—",
            _local_time(r["check_out_time"]) or "—",
            r["hours"] if r["hours"] is not None else "—",
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = BORDER
            if col_idx == 4:
                cell.fill = PatternFill("solid", fgColor=ATTENDANCE_STATUS_COLOR.get(r["status"], "FFFFFF"))
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")

    return _finish(wb)


ONBOARDING_HEADERS = ["Intern", "Intern code", "Department", "Designation", "Started", "Progress", "Status"]
ONBOARDING_WIDTHS = [22, 12, 18, 20, 12, 10, 14]
ONBOARDING_STATUS_COLOR = {"Completed": "C6EFCE", "On track": "FFEB9C", "Needs action": "FFC7CE"}


def build_onboarding_excel(rows, title):
    wb, ws = _sheet("Onboarding")
    _write_title_and_headers(ws, title, ONBOARDING_HEADERS, ONBOARDING_WIDTHS)

    for row_idx, r in enumerate(rows, start=3):
        values = [
            r["full_name"],
            r["intern_code"],
            r["department_name"] or "—",
            r["designation_name"] or "—",
            r["internship_start_date"],
            f"{r['progress_percent']}%",
            r["status_label"],
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = BORDER
            if col_idx == 7:
                cell.fill = PatternFill("solid", fgColor=ONBOARDING_STATUS_COLOR.get(r["status_label"], "FFFFFF"))
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")

    return _finish(wb)


EXIT_HEADERS = ["Employee", "Department", "Exit type", "Exit date", "Last day", "Exit interview", "Status"]
EXIT_WIDTHS = [22, 18, 14, 12, 12, 14, 14]
EXIT_STATUS_COLOR = {"Completed": "C6EFCE", "In progress": "FFEB9C", "Pending": "FFC7CE"}


def build_exit_excel(rows, title):
    wb, ws = _sheet("Exits")
    _write_title_and_headers(ws, title, EXIT_HEADERS, EXIT_WIDTHS)

    for row_idx, r in enumerate(rows, start=3):
        values = [
            r["full_name"],
            r["department_name"] or "—",
            r["exit_type_label"],
            r["exit_date"],
            r["last_working_date"] or "—",
            "Yes" if r["exit_interview_completed"] else "No",
            r["status_label"],
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = BORDER
            if col_idx == 7:
                cell.fill = PatternFill("solid", fgColor=EXIT_STATUS_COLOR.get(r["status_label"], "FFFFFF"))
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")

    return _finish(wb)
