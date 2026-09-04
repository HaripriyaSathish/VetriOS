from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADERS = ["S.No", "Date", "Trainer Name", "Trainee Name", "Status", "Course Name", "Topic Covered"]
COL_WIDTHS = [6, 12, 18, 18, 14, 22, 30]
STATUS_COLOR = {"Present": "C6EFCE", "Absent": "FFC7CE", "Not Marked": "F2F2F2"}


def build_training_log_excel(rows, title):
    wb = Workbook()
    ws = wb.active
    ws.title = "Training Log"
    n_cols = len(HEADERS)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.fill = PatternFill("solid", fgColor="B8CCE4")
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    thin = Side(style="thin", color="000000")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.fill = PatternFill("solid", fgColor="FFFF00")
        cell.font = Font(bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row_idx, r in enumerate(rows, start=3):
        values = [r["sno"], r["date"], r["trainer_name"], r["trainee_name"], r["status"], r["course_name"], r["topic"]]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border
            if col_idx == 5:
                cell.fill = PatternFill("solid", fgColor=STATUS_COLOR.get(r["status"], "FFFFFF"))
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")

    for col_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer